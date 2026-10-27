"""Vẽ giao diện các màn hình của Dev 2 (M2, M4 phần thách đấu và kết thúc ván, giao diện chính) trên Excel theo kiểu
giáo trình: khung màn hình viền đen, tiêu đề đậm ở giữa, ô nhập viền mảnh, nút nền xanh, bảng có tiêu đề cột in nghiêng,
mũi tên xám giữa các màn hình. Mỗi màn hình một trang tính trong docs/ui/giao-dien.xlsx; Excel (qua COM) xuất từng trang
ra PDF, PyMuPDF vẽ lại ở độ phân giải cao và cắt bỏ lề trắng thành docs/ui/png/<trang>.png.

    python build_ui.py            # dựng file Excel và xuất ảnh (cần Microsoft Excel)
    python build_ui.py --no-png   # chỉ dựng file Excel

Toạ độ tính theo ô (hàng, cột), bắt đầu từ 1; mỗi cột rộng COL_W, mỗi hàng cao ROW_H.
"""
import argparse
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
XLSX = HERE / "giao-dien.xlsx"
PNG = HERE / "png"

COL_W, ROW_H, NCOL = 2.6, 18, 43
FONT = "Arial"
LINE = "000000"
BUTTON_FILL = "9DC3E6"
HIGHLIGHT_FILL = "FFF2CC"
NOTE_COLOR = "595959"
ARROW_COLOR = "808080"


class Screen:
    def __init__(self, wb, name, last_row):
        self.ws = wb.create_sheet(name)
        self.ws.sheet_view.showGridLines = False
        for c in range(1, NCOL + 1):
            self.ws.column_dimensions[get_column_letter(c)].width = COL_W
        for r in range(1, last_row + 2):
            self.ws.row_dimensions[r].height = ROW_H
        self.ws.print_area = f"A1:{get_column_letter(NCOL)}{last_row + 1}"
        self.ws.page_setup.orientation = "landscape"
        self.ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.ws.page_setup.fitToWidth = self.ws.page_setup.fitToHeight = 1

    def _side(self, r, c, **sides):
        cell = self.ws.cell(r, c)
        b = cell.border
        cell.border = Border(**{k: sides.get(k, getattr(b, k)) for k in ("left", "right", "top", "bottom")})

    def outline(self, r1, c1, r2, c2, color=LINE, style="thin"):
        s = Side(style=style, color=color)
        for c in range(c1, c2 + 1):
            self._side(r1, c, top=s)
            self._side(r2, c, bottom=s)
        for r in range(r1, r2 + 1):
            self._side(r, c1, left=s)
            self._side(r, c2, right=s)

    def text(self, r, c1, c2, txt, size=10, bold=False, italic=False, color="000000", align="left", r2=None):
        r2 = r2 or r
        if (r2, c2) != (r, c1):
            self.ws.merge_cells(start_row=r, start_column=c1, end_row=r2, end_column=c2)
        cell = self.ws.cell(r, c1, txt)
        cell.font = Font(name=FONT, size=size, bold=bold, italic=italic, color=color)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=r2 > r)
        return cell

    def frame(self, r1, c1, r2, c2, title=None, user=None):
        """Khung một màn hình hoặc dialog: viền đen, tiêu đề đậm ở giữa, tên người chơi ở góc phải nếu có."""
        self.outline(r1, c1, r2, c2)
        if title:  # có tên người chơi ở góc phải thì chừa đều hai bên để tiêu đề vẫn ở giữa khung
            pad = 9 if user else 1
            self.text(r1 + 1, c1 + pad, c2 - pad, title, size=14, bold=True, align="center")
        if user:
            self.text(r1 + 1, c2 - 8, c2 - 1, user, align="right")

    def button(self, r, c1, c2, txt):
        cell = self.text(r, c1, c2, txt, align="center", r2=r + 1)
        for rr in (r, r + 1):
            for c in range(c1, c2 + 1):
                self.ws.cell(rr, c).fill = PatternFill("solid", fgColor=BUTTON_FILL)
        self.outline(r, c1, r + 1, c2)
        return cell

    def table(self, r, c, cols, rows, highlight=None):
        """cols: [(tiêu đề, số ô)]; rows: các hàng dữ liệu; highlight: chỉ số hàng được tô màu (người chơi đang xem)."""
        for i, row in enumerate([[h for h, _ in cols]] + rows):
            x = c
            for (_, w), val in zip(cols, row):
                cell = self.text(r + i, x, x + w - 1, val, bold=i == 0 or i - 1 == highlight, italic=i == 0)
                self.outline(r + i, x, r + i, x + w - 1)
                if i and i - 1 == highlight:
                    for cc in range(x, x + w):
                        self.ws.cell(r + i, cc).fill = PatternFill("solid", fgColor=HIGHLIGHT_FILL)
                x += w

    def arrow(self, r, c, label=None):
        self.text(r, c, c + 3, "⇨", size=30, color=ARROW_COLOR, align="center", r2=r + 1).font = Font(
            name="Segoe UI Symbol", size=30, color=ARROW_COLOR)
        if label:
            self.text(r - 1, c - 1, c + 4, label, size=9, italic=True, color=NOTE_COLOR, align="center")


def build():
    wb = Workbook()
    wb.remove(wb.active)

    # Hình 3.5: giao diện chính (khách thấy "Khách" thay cho tên người chơi)
    s = Screen(wb, "ui-main", 14)
    s.frame(2, 2, 14, 41, "PIKACHU", user="Người chơi A")
    s.button(6, 8, 19, "Chơi với máy")
    s.button(6, 24, 35, "Đấu đối kháng")
    s.button(10, 8, 19, "Xem bảng xếp hạng")
    s.button(10, 24, 35, "Cài đặt")

    # Hình 3.6: bảng xếp hạng, đang ở trang có dòng của người chơi (sau khi bấm Vị trí của tôi)
    s = Screen(wb, "ui-ranking", 15)
    s.frame(2, 2, 15, 41, "Bảng xếp hạng")
    s.table(5, 4, [("Hạng", 3), ("Tên", 10), ("Tổng điểm", 6), ("TB điểm đối thủ", 8), ("TB thời gian thắng", 9)],
            [["11", "Người chơi N", "4.5", "7.5", "07:45"], ["12", "Người chơi A", "3.5", "8.6", "06:58"]],
            highlight=1)
    s.button(9, 8, 14, "Trang trước")
    s.text(9, 16, 24, "Trang 2/2", align="center", r2=10)
    s.button(9, 26, 32, "Trang sau")
    s.button(12, 8, 14, "Quay lại")
    s.button(12, 26, 32, "Vị trí của tôi")

    # Hình 3.10: màn hình đấu đối kháng
    s = Screen(wb, "ui-lobby", 15)
    s.frame(2, 2, 15, 41, "Đấu đối kháng", user="Người chơi A")
    s.text(5, 4, 18, "Người chơi online", bold=True)
    s.text(5, 24, 38, "(click tên để thách đấu)", size=9, italic=True, color=NOTE_COLOR, align="right")
    s.table(6, 4, [("TT", 3), ("Tên", 14), ("Tổng điểm", 8), ("Trạng thái", 10)],
            [["1", "Người chơi B", "12.5", "Rỗi"], ["2", "Người chơi C", "10.5", "Bận"],
             ["3", "Người chơi D", "9", "Luyện tập"]])
    s.text(11, 4, 25, "B đã từ chối lời mời", color="0070C0")
    s.button(13, 4, 10, "Quay lại")

    # Hình 3.11: dialog xác nhận thách đấu, sau khi bấm Xác nhận đổi thành popup chờ
    s = Screen(wb, "ui-confirm", 10)
    s.frame(2, 2, 10, 17, "Xác nhận thách đấu")
    s.text(5, 3, 16, "Bạn muốn thách đấu B?", align="center")
    s.button(7, 4, 8, "Hủy")
    s.button(7, 11, 15, "Xác nhận")
    s.arrow(5, 19, "Xác nhận")
    s.frame(2, 25, 10, 41, "Đang chờ trả lời")
    s.text(5, 26, 40, "Đang chờ B trả lời... (còn 25 giây)", align="center")
    s.button(7, 29, 37, "Hủy lời mời")

    # Hình 3.12: popup lời mời trên máy người được mời
    s = Screen(wb, "ui-invitation", 13)
    s.frame(2, 2, 10, 22, "Thách đấu!")
    s.text(5, 3, 21, "A muốn thách đấu bạn (còn 28 giây)", align="center")
    s.button(7, 5, 10, "Reject")
    s.button(7, 14, 19, "Accept")
    s.text(12, 2, 30, "Popup không chặn, hiện ở mép trên màn hình; màn hình bên dưới vẫn thao tác được.",
           size=9, italic=True, color=NOTE_COLOR)
    s.text(13, 2, 30, "Nhiều lời mời thì hiện lần lượt.", size=9, italic=True, color=NOTE_COLOR)

    # Hình 3.14: dialog kết quả; bên phải là trường hợp đối thủ đã thoát
    s = Screen(wb, "ui-result", 11)
    s.frame(2, 2, 11, 20, "Chiến thắng!")
    s.text(5, 3, 19, "Bạn thắng (+1 điểm). Thời gian hoàn thành 06:42.", align="center")
    s.text(6, 3, 19, "Bạn có muốn tái đấu không? (còn 25 giây)", align="center")
    s.button(8, 5, 10, "Dừng")
    s.button(8, 13, 18, "Tái đấu")
    s.frame(2, 23, 11, 41, "Chiến thắng!")
    s.text(5, 24, 40, "Bạn thắng (+1 điểm). Đối thủ đã thoát.", align="center")
    s.text(6, 24, 40, "(còn 25 giây)", align="center")
    s.button(8, 28, 35, "Về sảnh")

    wb.save(XLSX)
    print("saved", XLSX)
    return wb.sheetnames


def export_png(names):
    """Excel xuất từng trang tính ra PDF (vẽ dạng vector), PyMuPDF vẽ lại ở 4x và cắt lề trắng."""
    import pymupdf
    import win32com.client
    from PIL import Image, ImageChops

    PNG.mkdir(exist_ok=True)
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        book = excel.Workbooks.Open(str(XLSX))
        with tempfile.TemporaryDirectory() as tmp:
            for name in names:
                pdf = Path(tmp) / f"{name}.pdf"
                book.Worksheets(name).ExportAsFixedFormat(0, str(pdf))
                with pymupdf.open(pdf) as doc:
                    pix = doc[0].get_pixmap(matrix=pymupdf.Matrix(4, 4))
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                box = ImageChops.difference(img, Image.new("RGB", img.size, "white")).getbbox()
                pad = 12
                img = img.crop((max(0, box[0] - pad), max(0, box[1] - pad),
                                min(img.width, box[2] + pad), min(img.height, box[3] + pad)))
                img.save(PNG / f"{name}.png")
                print("exported", name, img.size)
        book.Close(False)
    finally:
        excel.Quit()


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-png", action="store_true")
    names = build()
    if not ap.parse_args().no_png:
        export_png(names)


if __name__ == "__main__":
    main()
